"""Download a web page or PDF once and keep it: raw bytes, metadata and a normalised text copy.

    python -m bmw_feedback competitors fetch URL [--render]     -> prints the source id
    python -m bmw_feedback competitors grep SOURCE_ID REGEX     -> matches in the normalised text, to copy as evidence
    python -m bmw_feedback competitors sources                  -> every downloaded source

Each source lives in its own folder, competitor-data/raw/<source_id>/ (body.<ext> + meta.json), and its text in
competitor-data/text/<source_id>.txt. Downloaded files are data: they are parsed here, never executed. The network
is used only by `fetch`; the build and the app read the saved files.
"""

import hashlib
import html
import json
import re
import time
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

from .registry import RAW_DIR, TEXT_DIR

UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/128.0.0.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/pdf,*/*;q=0.8",
           "Accept-Language": "en-GB,en;q=0.9,de;q=0.8,zh-CN;q=0.6"}
BLOCK_HINTS = ("access denied", "just a moment", "attention required", "captcha", "are you a robot",
               "enable javascript and cookies", "request unsuccessful", "pardon our interruption")
SKIP_TAGS = {"script", "style", "noscript", "svg", "template", "iframe", "canvas"}
BLOCK_TAGS = {"p", "div", "section", "article", "header", "footer", "li", "ul", "ol", "tr", "table", "h1", "h2", "h3",
              "h4", "h5", "h6", "br", "dt", "dd", "dl", "blockquote", "figcaption", "main", "aside", "nav", "form",
              "caption", "thead", "tbody", "tfoot", "summary", "details", "option", "label", "hr"}
CELL_TAGS = {"td", "th"}


def source_id(url: str) -> str:
    return "s" + hashlib.sha1(url.strip().encode()).hexdigest()[:10]


def publisher(url: str) -> str:
    host = urlparse(url).netloc.lower()
    return host[4:] if host.startswith("www.") else host


# ---------------------------------------------------------------------- normalisation
_SWAP = {" ": " ", " ": " ", " ": " ", " ": " ", "​": "", "–": "-", "—": "-",
         "−": "-", "‐": "-", "‑": "-", "‘": "'", "’": "'", "‚": "'", "“": '"',
         "”": '"', "„": '"', "´": "'", "′": "'", "″": '"', "…": "...", "﻿": ""}
_SWAP_RE = re.compile("|".join(map(re.escape, _SWAP)))


def normalise(text: str) -> str:
    """The canonical form evidence is compared in: NFKC, plain quotes and dashes, single spaces."""
    text = unicodedata.normalize("NFKC", text or "")
    text = _SWAP_RE.sub(lambda m: _SWAP[m.group(0)], text)
    return re.sub(r"[ \t\r\f\v]+", " ", text)


def flat(text: str) -> str:
    """normalise() and also fold line breaks: what evidence matching uses (case is folded by the caller)."""
    return re.sub(r"\s+", " ", normalise(text)).strip()


class _Text(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.title: list[str] = []
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in SKIP_TAGS:
            self.skip += 1
        elif tag == "title":
            self.in_title = True
        elif tag in BLOCK_TAGS:
            self.out.append("\n")
        elif tag in CELL_TAGS:
            self.out.append(" | ")
        elif tag == "img" and not self.skip:
            alt = dict(attrs).get("alt")
            if alt:
                self.out.append(f" {alt} ")

    def handle_startendtag(self, tag, attrs):
        if tag == "br":
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
        elif tag == "title":
            self.in_title = False
        elif tag in BLOCK_TAGS:
            self.out.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
        if not self.skip:
            self.out.append(data)


def html_to_text(raw: str) -> tuple[str, str]:
    p = _Text()
    p.feed(raw)
    p.close()
    text = normalise(html.unescape("".join(p.out)))
    lines = [ln.strip() for ln in text.split("\n")]
    out, blank = [], False
    for ln in lines:
        if ln:
            out.append(ln)
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip(), flat(" ".join(p.title))


def pdf_to_text(data: bytes) -> tuple[str, str]:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    pages = [f"[[page {i}]]\n{normalise(page.get_text())}" for i, page in enumerate(doc, start=1)]
    title = (doc.metadata or {}).get("title") or ""
    return "\n".join(pages), flat(title)


def xlsx_to_text(data: bytes) -> tuple[str, str]:
    """Every sheet as lines of cells joined by ' | ' (values only, formulas evaluated by the writer)."""
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[[sheet {ws.title}]]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else (f"{v:g}" if isinstance(v, float) else str(v)) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                out.append(" | ".join(normalise(c).replace("\n", " ") for c in cells))
    return "\n".join(out), flat(wb.properties.title or "")


def json_to_text(data: bytes) -> tuple[str, str]:
    return normalise(json.dumps(json.loads(data), ensure_ascii=False, indent=1)), ""


# ---------------------------------------------------------------------- fetch
def _download(url: str, attempts: int = 4) -> tuple[bytes, str, str, int]:
    import httpx
    last = None
    for attempt in range(attempts):
        try:
            with httpx.Client(headers=HEADERS, follow_redirects=True, timeout=60.0, http2=False) as c:
                r = c.get(url)
        except httpx.HTTPError as e:            # transient network errors: retry with backoff
            last = e
            time.sleep(1.5 * (attempt + 1))
            continue
        if r.status_code in (429, 503) and attempt < attempts - 1:      # rate limited: wait as asked, then retry
            retry = r.headers.get("retry-after", "")
            time.sleep(min(float(retry) if retry.isdigit() else 8.0 * (attempt + 1), 60.0))
            continue
        return r.content, r.headers.get("content-type", ""), str(r.url), r.status_code
    raise RuntimeError(f"download failed for {url}: {last}")


def _render(url: str, wait_ms: int = 2500) -> tuple[bytes, str, str, int]:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as e:
        raise RuntimeError("--render needs Playwright: pip install playwright && python -m playwright install chromium") from e
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            page = browser.new_page(user_agent=UA, locale="en-GB", viewport={"width": 1366, "height": 2000})
            resp = page.goto(url, wait_until="domcontentloaded", timeout=60_000)
            try:
                page.wait_for_load_state("networkidle", timeout=20_000)
            except Exception:                    # some pages never go idle; what has rendered is kept
                pass
            page.wait_for_timeout(wait_ms)
            body = page.content().encode("utf-8")
            status = resp.status if resp else 0
            return body, "text/html; charset=utf-8", page.url, status
        finally:
            browser.close()


def _kind(content_type: str, body: bytes, url: str) -> str:
    ct = content_type.lower()
    if "pdf" in ct or body[:5] == b"%PDF-" or urlparse(url).path.lower().endswith(".pdf"):
        return "pdf"
    if "json" in ct:
        return "json"
    if "spreadsheetml" in ct or body[:2] == b"PK" and urlparse(url).path.lower().endswith((".xlsx", ".xlsm")):
        return "xlsx"
    return "html"


def _decode(body: bytes, content_type: str) -> str:
    m = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    for enc in ([m.group(1)] if m else []) + ["utf-8", "gb18030", "latin-1"]:
        try:
            return body.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return body.decode("utf-8", errors="replace")


def fetch(url: str, render: bool = False, force: bool = False) -> dict:
    """Download (or render) a URL into its own folder and write the normalised text. Returns the metadata."""
    sid = source_id(url)
    folder = RAW_DIR / sid
    meta_path = folder / "meta.json"
    if meta_path.exists() and not force:
        meta = json.loads(meta_path.read_text())
        if meta.get("ok"):
            return {**meta, "cached": True}
    body, ctype, final_url, status = (_render if render else _download)(url)
    kind = _kind(ctype, body, final_url)
    try:
        if kind == "pdf":
            text, title = pdf_to_text(body)
        elif kind == "json":
            text, title = json_to_text(body)
        elif kind == "xlsx":
            text, title = xlsx_to_text(body)
        else:
            text, title = html_to_text(_decode(body, ctype))
    except Exception as e:                       # unparseable payload: keep it, report it, never pretend
        text, title = "", f"(text extraction failed: {e})"
    low = text[:4000].lower()
    blocked = status >= 400 or len(text) < 400 or any(h in low for h in BLOCK_HINTS)
    folder.mkdir(parents=True, exist_ok=True)
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    (folder / f"body.{kind}").write_bytes(body)
    (TEXT_DIR / f"{sid}.txt").write_text(text)
    meta = {"source_id": sid, "url": url, "final_url": final_url, "publisher": publisher(final_url or url),
            "status": status, "content_type": ctype, "kind": kind, "rendered": render, "title": title,
            "retrieved_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "bytes": len(body),
            "sha256": hashlib.sha256(body).hexdigest(), "text_chars": len(text), "ok": not blocked,
            "problem": None if not blocked else ("HTTP status %d" % status if status >= 400 else
                                                 "page looks blocked or empty (try --render)")}
    meta_path.write_text(json.dumps(meta, indent=1, ensure_ascii=False))
    return meta


# ---------------------------------------------------------------------- read
def meta(sid: str) -> dict | None:
    p = RAW_DIR / sid / "meta.json"
    return json.loads(p.read_text()) if p.exists() else None


def text(sid: str) -> str | None:
    p = TEXT_DIR / f"{sid}.txt"
    return p.read_text() if p.exists() else None


def all_sources() -> list[dict]:
    if not RAW_DIR.exists():
        return []
    return [json.loads(p.read_text()) for p in sorted(RAW_DIR.glob("*/meta.json"))]


def grep(sid: str, pattern: str, context: int = 160, limit: int = 25) -> list[str]:
    """Matches of a regex (case-insensitive) in the flat normalised text, with context, ready to copy."""
    t = text(sid)
    if t is None:
        raise FileNotFoundError(f"no text for source {sid}; fetch it first")
    f = flat(t)
    out = []
    for m in re.finditer(pattern, f, re.I):
        a, b = max(0, m.start() - context), min(len(f), m.end() + context)
        out.append(f[a:b])
        if len(out) >= limit:
            break
    return out


def path_for(sid: str) -> Path:
    return TEXT_DIR / f"{sid}.txt"
