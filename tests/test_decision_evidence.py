"""Evidence release gates use subsets of the private, real source workbooks."""
from dataclasses import replace

import pandas as pd
import pytest

from pm_copilot.decision_evidence import build_evidence, validate_citations


def _context(data, **changes):
    row = data.findings.iloc[0]
    return {"vehicle": row["derivative"], "markets": [row["region"]],
            "launch_year": 2030, **changes}


def test_A1_duplicate_classifications_do_not_multiply_reviews(real_feedback_subset):
    data = real_feedback_subset
    component = data.findings.iloc[0]["component"]
    context = _context(data)
    before = build_evidence(context, [f"component: {component}"], data)
    repeated = replace(data, findings=pd.concat([data.findings, data.findings], ignore_index=True))
    after = build_evidence(context, [f"component: {component}"], repeated)
    assert before["features"][0]["markets"] == after["features"][0]["markets"]


def test_A2_frozen_raw_inputs_produce_identical_evidence(real_feedback_subset):
    data = real_feedback_subset
    component = data.findings.iloc[0]["component"]
    first = build_evidence(_context(data), [f"component: {component}"], data)
    second = build_evidence(_context(data), [f"component: {component}"], data)
    assert first == second
    assert first["source_hashes"]


def test_A8_all_counted_evidence_has_resolvable_original_locators(real_feedback_subset):
    data = real_feedback_subset
    component = data.findings.iloc[0]["component"]
    pack = build_evidence(_context(data), [f"component: {component}"], data)
    region = next(iter(pack["features"][0]["markets"].values()))
    assert region["sample"]["n"] > 0
    validate_citations(pack, region["direct_evidence_ids"])
    for citation_id in region["direct_evidence_ids"]:
        cite = pack["citations"][citation_id]
        assert cite["file"].endswith(".xlsx")
        assert cite["sheet"] == "Feedback_Explorer"
        assert cite["row"] >= 2
        assert cite["raw_record_id"]
        assert cite["file"] in pack["source_hashes"]
    with pytest.raises(ValueError, match="citation"):
        validate_citations(pack, ["not-an-evidence-id"])


def test_A9_subfeature_does_not_fall_back_to_parent_component(real_bmw_data):
    from pm_copilot import taxonomy as tx

    feature = next(f for f in tx.FEATURES if f.key == "seat_ventilation")
    rows = real_bmw_data.findings
    parent = rows[rows["component"].isin(feature.components)]
    unrelated = parent[~parent["text"].str.contains(feature.component_text_filter, case=False, regex=True)]
    if unrelated.empty:
        pytest.skip("Real workbook has no parent-only examples for this feature.")
    subset = replace(real_bmw_data, findings=unrelated.copy())
    pack = build_evidence(_context(subset), [feature.key], subset)
    region = next(iter(pack["features"][0]["markets"].values()))
    assert region["counts"]["mentions"] == 0
    assert region["counts"]["complaints"] == 0
    assert region["contextual_evidence_ids"]


def test_A8_powertrain_and_market_share_the_same_denominator(real_feedback_subset):
    data = real_feedback_subset
    row = data.findings.iloc[0]
    context = _context(data, powertrain=row["powertrain"])
    pack = build_evidence(context, [f"component: {row['component']}"], data)
    region = pack["features"][0]["markets"][row["region"]]
    scoped = data.findings[data.findings["derivative"].eq(row["derivative"])
                           & data.findings["region"].eq(row["region"])
                           & data.findings["powertrain"].eq(row["powertrain"])]
    assert region["sample"]["n"] == scoped["verbatim_id"].nunique()
    assert set(region["sample"]["powertrain_mix"]) == {row["powertrain"]}
    assert all(k.startswith("Source ") or k == "Unknown source" for k in region["sample"]["source_mix"])


def test_A9_unreviewed_model_labels_do_not_count_as_observations(real_feedback_subset):
    data = real_feedback_subset
    component = data.findings.iloc[0]["component"]
    changed = data.findings.copy()
    changed["label_source"] = "model"
    pack = build_evidence(_context(data), [f"component: {component}"], replace(data, findings=changed))
    region = next(iter(pack["features"][0]["markets"].values()))
    assert region["counts"]["mentions"] == 0
    assert region["unreviewed_evidence_ids"]


def test_A9_unavailable_launch_year_sales_are_unknown(real_feedback_subset):
    data = real_feedback_subset
    component = data.findings.iloc[0]["component"]
    pack = build_evidence(_context(data, launch_year=2029), [f"component: {component}"], data)
    region = next(iter(pack["features"][0]["markets"].values()))
    assert region["sales"]["volume"] is None
    assert region["sales"]["status"] == "unknown"


def test_A8_survey_preserves_original_cells_and_scale(real_bmw_data):
    pack = build_evidence({"vehicle": "G60", "markets": ["EU", "CN", "US"], "launch_year": 2030},
                          ["seat_comfort"], real_bmw_data)
    surveys = [s for market in pack["features"][0]["markets"].values() for s in market["survey"]]
    assert surveys
    for survey in surveys:
        assert survey["original_scale"]
        assert survey["source_ids"]
        for cid in survey["source_ids"]:
            assert pack["citations"][cid]["cell"]
    us = [s for s in surveys if s["study"] == "US_2025"]
    assert us and us[0]["distribution"]


def test_A9_observed_zero_volume_is_distinct_from_missing(real_bmw_data):
    zeros = real_bmw_data.sales[real_bmw_data.sales["volume_2030"].eq(0)]
    if zeros.empty:
        pytest.skip("No zero-volume market exists in the supplied real sales workbook.")
    row = zeros.iloc[0]
    context = {"vehicle": row["derivative"], "markets": [row["region"]], "launch_year": 2030}
    pack = build_evidence(context, ["seat_comfort"], real_bmw_data)
    sales = pack["features"][0]["markets"][row["region"]]["sales"]
    assert sales["volume"] == 0
    assert sales["status"] == "available"
    assert sales["label"] == "scenario evaluation"
    assert sales["basis"] == "BMW projection"
    assert sales["source_ids"]


def test_A8_feedback_excerpt_matches_original_source_cell(real_feedback_subset):
    from openpyxl import load_workbook
    from pm_copilot.config import RAW_DIR

    data = real_feedback_subset
    pack = build_evidence(_context(data), [f"component: {data.findings.iloc[0]['component']}"], data)
    region = next(iter(pack["features"][0]["markets"].values()))
    citation = pack["citations"][region["direct_evidence_ids"][0]]
    workbook = load_workbook(RAW_DIR / citation["file"], read_only=True, data_only=True)
    try:
        sheet = workbook[citation["sheet"]]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        text_column = headers.index("Customer Feedback") + 1
        assert citation["text"] == str(sheet.cell(citation["row"], text_column).value)
    finally:
        workbook.close()


def test_A2_runtime_read_does_not_write_processed_artifacts(real_bmw_data):
    from pm_copilot.config import PROCESSED_DIR
    from pm_copilot.store import load_runtime

    before = {p.name: (p.stat().st_size, p.stat().st_mtime_ns) for p in PROCESSED_DIR.iterdir() if p.is_file()}
    runtime = load_runtime()
    after = {p.name: (p.stat().st_size, p.stat().st_mtime_ns) for p in PROCESSED_DIR.iterdir() if p.is_file()}
    assert before == after
    assert runtime.manifest["sources"] == real_bmw_data.manifest["sources"]


def test_A2_feature_queries_preserve_alias_bindings_after_sort_and_dedup(real_feedback_subset):
    queries = ["seat_heating", "Heated steering wheel", "heated_steering_wheel"]
    first = build_evidence(_context(real_feedback_subset), queries, real_feedback_subset)
    second = build_evidence(_context(real_feedback_subset), list(reversed(queries)), real_feedback_subset)
    assert first == second
    assert first["feature_queries"] == {
        "seat_heating": "seat_heating", "Heated steering wheel": "heated_steering_wheel",
        "heated_steering_wheel": "heated_steering_wheel"}
    assert len(first["features"]) == 2
